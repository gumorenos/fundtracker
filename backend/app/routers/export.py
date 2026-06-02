from datetime import datetime
from io import BytesIO

import openpyxl
from fastapi import APIRouter, Depends, Query
from fastapi.responses import StreamingResponse
from sqlalchemy import func
from sqlalchemy.orm import Session

from app.auth import get_current_user
from app.database import get_db
from app.models import Transaction, User

router = APIRouter(prefix="/export", tags=["export"])

TYPE_LABELS = {
    "expense": "Gasto soles",
    "currency_exchange": "Cambio moneda",
    "usd_expense": "Gasto USD",
}

HEADER_STYLE = {
    "font": openpyxl.styles.Font(bold=True),
    "fill": openpyxl.styles.PatternFill("solid", fgColor="1E293B"),
}


def _style_header(ws, row=1):
    font = openpyxl.styles.Font(bold=True, color="FFFFFF")
    fill = openpyxl.styles.PatternFill("solid", fgColor="334155")
    for cell in ws[row]:
        cell.font = font
        cell.fill = fill


@router.get("/excel")
def export_excel(
    date_from: datetime | None = Query(None),
    date_to: datetime | None = Query(None),
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    q = db.query(Transaction).filter(Transaction.user_id == current_user.id)
    if date_from:
        q = q.filter(Transaction.transaction_date >= date_from)
    if date_to:
        q = q.filter(Transaction.transaction_date <= date_to)
    transactions = q.order_by(Transaction.transaction_date.desc()).all()

    wb = openpyxl.Workbook()

    # ── Sheet 1: Transactions ──────────────────────────────────────────────
    ws1 = wb.active
    ws1.title = "Transacciones"
    ws1.append([
        "Fecha", "Tipo", "Descripción", "Categoría", "Fondo",
        "USD", "PEN", "TC", "Notas", "Etiquetas",
    ])
    _style_header(ws1)
    for tx in transactions:
        ws1.append([
            tx.transaction_date.strftime("%Y-%m-%d"),
            TYPE_LABELS.get(tx.type, tx.type),
            tx.description or "",
            tx.category.name if tx.category else "",
            tx.fund.name if tx.fund else "",
            float(tx.amount_usd) if tx.amount_usd is not None else "",
            float(tx.amount_pen) if tx.amount_pen is not None else "",
            float(tx.exchange_rate) if tx.exchange_rate is not None else "",
            tx.notes or "",
            ", ".join(tx.tags) if tx.tags else "",
        ])
    for col in ws1.columns:
        ws1.column_dimensions[col[0].column_letter].width = 16

    # ── Sheet 2: Summary ──────────────────────────────────────────────────
    ws2 = wb.create_sheet("Resumen")

    # By category
    ws2.append(["Por categoría", "", ""])
    ws2.append(["Categoría", "Total PEN", "# Transacciones"])
    _style_header(ws2, row=2)
    cat_totals: dict[str, list] = {}
    for tx in transactions:
        if tx.type == "expense":
            name = tx.category.name if tx.category else "Sin categoría"
            if name not in cat_totals:
                cat_totals[name] = [0.0, 0]
            cat_totals[name][0] += float(tx.amount_pen or 0)
            cat_totals[name][1] += 1
    for name, (total, count) in sorted(cat_totals.items(), key=lambda x: -x[1][0]):
        ws2.append([name, round(total, 2), count])

    ws2.append([])

    # By type
    ws2.append(["Por tipo", "", ""])
    ws2.append(["Tipo", "Total USD", "Total PEN"])
    _style_header(ws2, row=ws2.max_row)
    type_totals: dict[str, list] = {}
    for tx in transactions:
        label = TYPE_LABELS.get(tx.type, tx.type)
        if label not in type_totals:
            type_totals[label] = [0.0, 0.0]
        type_totals[label][0] += float(tx.amount_usd or 0)
        type_totals[label][1] += float(tx.amount_pen or 0)
    for label, (usd, pen) in type_totals.items():
        ws2.append([label, round(usd, 2), round(pen, 2)])

    ws2.append([])

    # TC stats
    tc_values = [float(tx.exchange_rate) for tx in transactions if tx.exchange_rate]
    if tc_values:
        ws2.append(["Tipos de cambio", "", ""])
        ws2.append(["Promedio", "Mínimo", "Máximo"])
        _style_header(ws2, row=ws2.max_row)
        ws2.append([
            round(sum(tc_values) / len(tc_values), 4),
            round(min(tc_values), 4),
            round(max(tc_values), 4),
        ])

    for col in ws2.columns:
        ws2.column_dimensions[col[0].column_letter].width = 20

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)

    filename = f"fundtracker_{datetime.now().strftime('%Y%m%d_%H%M')}.xlsx"
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )
